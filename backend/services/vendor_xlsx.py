"""Vendor price list in and out of Excel — the way a refreshed price reaches an installed copy.

The PDF import needs the supplier's parse scripts, which an installed copy does not have
(issue #36). This is the same job without them: export the merged price table, edit it in Excel,
import it back. Changes are written to the vendor overlay (vendor_overlay.py), never to
lookup.json, so they survive a build_lookup.py re-run exactly like a hand correction (ADR-0001).

A blank cell means "no change" here, not "clear" — clearing an opinion is the vendor tab's own
"ล้างค่า". Two steps like the catalogue import: preview() writes only a staging file and lists
what would change, apply() backs the overlay up and writes.
"""

import io
import json
import re
import shutil
import time
import uuid

from backend import config
from backend.services import catalog, vendor_admin, vendor_lookup, vendor_overlay
from backend.validation import ValidationError

XLSX_NAME = "vendor-prices.xlsx"
SHEET = "ราคา"
COLUMNS = [
    ("รหัส", "code"),
    ("ชื่อ", "name"),
    ("ราคา", "price"),
    ("ขนาด (มม.)", "size_mm"),
    ("จำนวนต่อแพ็ก", "pack_qty"),
    ("หน่วยแพ็ก", "pack_unit"),
]
FIELDS = [field for _label, field in COLUMNS[1:]]
LISTED = 300  # how many changed rows the preview spells out; the count is always the full one


def _staging():
    return config.DATA_DIR / "import-staging"


def export_xlsx(supplier):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    if supplier not in config.VENDOR_SUPPLIERS:
        raise ValidationError(f"ไม่รู้จักซัพพลายเออร์ '{supplier}'")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append([label for label, _field in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for code, entry in sorted(vendor_lookup.entries().items()):
        if entry.get("supplier") != supplier:
            continue
        sheet.append([code] + [entry.get(field) for field in FIELDS])
    for letter, width in zip("ABCDEF", (16, 44, 10, 12, 14, 12)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _text(value):
    return "" if value is None else str(value).strip()


def _number(raw, label, where, errors, integer=False):
    text = _text(raw).replace(",", "")
    if not text:
        return None
    try:
        value = float(text)
        if value < 0 or (integer and (value != int(value) or value < 1)):
            raise ValueError
    except ValueError:
        errors.append(f"{where}: {label} '{text}' ต้องเป็น{'จำนวนเต็มตั้งแต่ 1' if integer else 'ตัวเลขตั้งแต่ 0'}ขึ้นไป")
        return None
    return int(value) if integer else value


def preview(file, supplier):
    from openpyxl import load_workbook

    if supplier not in config.VENDOR_SUPPLIERS:
        raise ValidationError(f"ไม่รู้จักซัพพลายเออร์ '{supplier}'")
    try:
        sheet = load_workbook(file, read_only=True, data_only=True)[SHEET]
    except Exception:
        raise ValidationError(f"เปิดไฟล์ไม่ได้ — ใช้ไฟล์ .xlsx ที่ส่งออกจากหน้านี้ (ชีต '{SHEET}')") from None
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValidationError("ไฟล์ว่าง")
    header = [_text(cell) for cell in rows[0]]
    labels = [label for label, _field in COLUMNS]
    if any(label not in header for label in labels):
        raise ValidationError("หัวคอลัมน์ไม่ตรง — ใช้ไฟล์ที่ส่งออกจากหน้านี้ ไม่เปลี่ยนชื่อหัวคอลัมน์")
    column = {field: header.index(label) for label, field in COLUMNS}

    entries = vendor_lookup.entries()
    errors, warnings, changed, seen = [], [], [], set()
    unchanged = 0
    for number, cells in enumerate(rows[1:], start=2):
        get = lambda field: cells[column[field]] if column[field] < len(cells) else None
        code = _text(get("code"))
        if not code:
            continue
        where = f"แถว {number} ({code})"
        if code in seen:
            errors.append(f"{where}: รหัสซ้ำในไฟล์")
            continue
        seen.add(code)
        if code not in entries:
            try:
                catalog.find(code)
            except ValidationError:
                errors.append(f"{where}: ไม่รู้จักรหัสนี้ (ไม่มีทั้งในราคาซัพพลายเออร์และ catalogue)")
                continue

        new = {
            "name": _text(get("name")) or None,
            "price": _number(get("price"), "ราคา", where, errors),
            "size_mm": _number(get("size_mm"), "ขนาด", where, errors),
            "pack_qty": _number(get("pack_qty"), "จำนวนต่อแพ็ก", where, errors, integer=True),
            "pack_unit": _text(get("pack_unit")) or None,
        }
        current = entries.get(code, {})
        diff = {}
        for field, value in new.items():
            if value is None or value == current.get(field):
                continue
            if field == "size_mm" and vendor_admin._catalog_size_mm(code) is not None:
                warnings.append(f"{where}: ข้ามขนาด — ใช้ขนาดจาก catalogue อยู่แล้ว")
                continue
            diff[field] = [current.get(field), value]
        if diff:
            changed.append({"code": code, "changes": diff})
        else:
            unchanged += 1

    token = uuid.uuid4().hex
    _staging().mkdir(parents=True, exist_ok=True)
    (_staging() / f"vendor-{token}.json").write_text(
        json.dumps(changed, ensure_ascii=False), encoding="utf-8"
    )
    return {
        "token": token,
        "rows": len(seen),
        "changed": len(changed),
        "unchanged": unchanged,
        "listed": changed[:LISTED],
        "errors": errors,
        "warnings": warnings,
        "can_apply": not errors and bool(changed),
    }


def apply(token):
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise ValidationError("ไม่รู้จักการนำเข้านี้")
    path = _staging() / f"vendor-{token}.json"
    if not path.is_file():
        raise ValidationError("ไม่พบข้อมูลที่ตรวจไว้ — อัปโหลดไฟล์ใหม่อีกครั้ง")
    changed = json.loads(path.read_text(encoding="utf-8"))

    backup = None
    overlay = vendor_overlay.overlay_path()
    if overlay.is_file():
        backups = config.DATA_DIR / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        backup = f"vendor-overlay-{time.strftime('%Y%m%d-%H%M%S')}.json"
        shutil.copyfile(overlay, backups / backup)

    fields = set()
    for item in changed:
        fields.update(item["changes"])
    entries = vendor_lookup.entries()
    writes = {}
    for item in changed:
        code = item["code"]
        # every field this import speaks for is restated per code (set_many drops the rest), so
        # a field this code did not change carries its current merged value forward
        wanted = {
            field: item["changes"][field][1] if field in item["changes"] else entries.get(code, {}).get(field)
            for field in fields
        }
        # what the supplier's own sheet already says is no opinion at all
        writes[code] = {
            field: value for field, value in wanted.items()
            if value is not None and value != vendor_lookup.base_field(code, field)
        }
    vendor_overlay.set_many(writes, speaks_for=tuple(fields))
    vendor_lookup.refresh()
    path.unlink(missing_ok=True)
    return {"changed": len(changed), "backup": backup}

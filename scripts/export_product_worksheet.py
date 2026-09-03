"""Staff data-entry worksheet: every catalogue product, one row each, with a thumbnail and
every field the app actually needs — highlighting whichever of them is still blank.

Ad-hoc export, not part of the app or its tests — run it whenever the catalogue needs a fresh
"fill this in" pass (started for wayfinder ticket #2, backfilling `price` for the budget-wizard
map at issue #1). Needs openpyxl (`pip install openpyxl` into .venv — not in requirements.txt,
this script never runs in production).

    python scripts/export_product_worksheet.py
"""

import json

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from _bootstrap import ROOT

from backend.services import catalog

OUT_PATH = ROOT / "catalog" / "product_worksheet.xlsx"
THUMB_PX = 60

RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFE8A3")
HEADER_FILL = PatternFill("solid", fgColor="2E4A3B")
HEADER_FONT = Font(bold=True, color="FFFFFF")

COLUMNS = ["#", "รหัสสินค้า", "รูปสินค้า", "หมวด", "ขนาด (ตามแคตตาล็อก)", "ขนาด (มม.)",
           "ทรง", "สี", "ราคา (บาท)", "หมายเหตุ — พนักงานกรอกเพิ่ม"]


def load_descriptions():
    path = catalog.config.CATALOG_PATH.parent / "descriptions.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}


def siblings_by_image(rows):
    """code -> every code (including itself) whose photo is byte-identical to it.

    Every code gets its own <code>.png filename even when several codes were printed with
    the exact same photo (catalog.py's crop_is_ambiguous docstring: "35072-1.png is
    byte-identical to 34072-1, 36072-1..."), so grouping by filename finds nothing — this
    groups by the file's actual content hash instead, the same signal
    product_images.json's own `shared_with` count was built from.
    """
    import hashlib

    digest_of = {}
    by_digest = {}
    for row in rows:
        path = catalog.image_path(row["code"])
        if not path or not path.is_file():
            continue
        digest = hashlib.sha1(path.read_bytes()).hexdigest()
        digest_of[row["code"]] = digest
        by_digest.setdefault(digest, []).append(row["code"])
    return {code: by_digest[digest] for code, digest in digest_of.items()}


def hard_image_problem(code):
    """Reasons a photo must never be shown at all, not even with a caveat — it is not a
    picture of any real product (page furniture) or the code itself is in dispute. Distinct
    from crop_is_ambiguous(), which is a *shared* photo of real products — see
    ambiguous_note() below for why that one still gets shown.
    """
    if not catalog.image_for(code):
        return "ไม่มีรูป"
    if catalog.crop_is_not_a_product(code):
        return f"ไม่ใช่รูปสินค้า ({catalog._crop_verdicts().get(code)})"
    if catalog.code_is_contested(code):
        return "โค้ดนี้มีสองความหมายในแคตตาล็อก — ดูหน้า Settings"
    return None


def ambiguous_note(code, groups):
    """When this code's photo is shared with others: which ones, so staff can at least tell
    "this is one of these N products" from a real picture, rather than nothing. None when the
    photo is this code's alone."""
    if not catalog.crop_is_ambiguous(code):
        return None
    others = [c for c in groups.get(code, []) if c != code]
    return f"รูปนี้ใช้ร่วมกับ {', '.join(others)} — เช็ครหัสให้ตรงก่อนยืนยันสินค้าจริง"


def thumbnail_bytes(path):
    """A small square PNG for the sheet, or None if the crop is missing/unreadable."""
    if not path or not path.is_file():
        return None
    try:
        img = PILImage.open(path).convert("RGBA")
    except Exception:
        return None
    img.thumbnail((THUMB_PX, THUMB_PX))
    canvas = PILImage.new("RGBA", (THUMB_PX, THUMB_PX), (255, 255, 255, 0))
    canvas.paste(img, ((THUMB_PX - img.width) // 2, (THUMB_PX - img.height) // 2), img)
    import io

    buf = io.BytesIO()
    canvas.save(buf, format="PNG")
    buf.seek(0)
    return buf


def build():
    rows = json.loads(catalog.config.CATALOG_PATH.read_text(encoding="utf-8"))
    descriptions = load_descriptions()
    groups = siblings_by_image(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    widths = [4, 12, 26, 16, 16, 11, 12, 12, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    missing_image = missing_size = missing_shape = missing_colour = missing_price = ambiguous_count = 0

    for i, row in enumerate(rows, start=1):
        r = i + 1
        ws.row_dimensions[r].height = THUMB_PX * 0.78  # points ≈ px * 0.75, a hair extra for padding

        code = row["code"]
        attrs = (descriptions.get(code) or {}).get("attributes") or {}
        mm = catalog.longest_side_mm(row)
        category = catalog.category_of(row)
        shape = attrs.get("shape") or ""
        colour = attrs.get("primary_colour") or ""
        price = row.get("price")

        ws.cell(r, 1, i)
        ws.cell(r, 2, code)

        problem = hard_image_problem(code)
        thumb = None if problem else thumbnail_bytes(catalog.image_path(code))
        if thumb:
            xl_img = XLImage(thumb)
            xl_img.width = xl_img.height = THUMB_PX
            ws.add_image(xl_img, f"C{r}")
        else:
            ws.cell(r, 3, problem or "รูปอ่านไฟล์ไม่ได้").fill = RED
            missing_image += 1

        note = ambiguous_note(code, groups)
        if note:
            ws.cell(r, 10, note).fill = AMBER
            ambiguous_count += 1

        ws.cell(r, 4, category or "")
        if not category:
            ws.cell(r, 4).fill = RED

        ws.cell(r, 5, row.get("size_raw") or "")
        size_cell = ws.cell(r, 6, mm if mm is not None else None)
        if mm is None:
            ws.cell(r, 5).fill = RED
            size_cell.fill = RED
            missing_size += 1

        shape_cell = ws.cell(r, 7, shape)
        if not shape:
            shape_cell.fill = RED
            missing_shape += 1

        colour_cell = ws.cell(r, 8, colour)
        if not colour:
            colour_cell.fill = RED
            missing_colour += 1

        price_cell = ws.cell(r, 9, price)
        if price is None:
            price_cell.fill = RED
            missing_price += 1

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH} — {len(rows)} products")
    print(f"missing: image={missing_image} size={missing_size} shape={missing_shape} "
          f"colour={missing_colour} price={missing_price}")
    print(f"ambiguous (photo shown, sibling note added): {ambiguous_count}")


if __name__ == "__main__":
    build()

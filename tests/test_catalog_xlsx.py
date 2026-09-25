"""The catalogue as products.xlsx + images/ in a zip (ADR-0005): export, preview, apply.

Import replaces every product except MS Natural Design, only after a preview with no row
errors, and backs the old catalogue up first."""

import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

from backend import config
from backend.services import (
    catalog, catalog_admin, catalog_xlsx, settings, shop_overlay, vendor_lookup,
)
from backend.validation import ValidationError

from helpers import png_bytes

BC = "Bangkok Christmas"
MSND = "MS Natural Design"


@pytest.fixture
def temp_catalog(tmp_path, monkeypatch):
    folder = tmp_path / "catalog"
    folder.mkdir()
    products = folder / "products.json"
    products.write_text("[]", encoding="utf-8")
    monkeypatch.setattr(catalog_admin, "PRODUCTS_PATH", products)
    monkeypatch.setattr(catalog_admin, "IMAGES_DIR", folder / "images")
    monkeypatch.setattr(catalog_admin, "PRODUCT_IMAGES_PATH", folder / "product_images.json")
    monkeypatch.setattr(config, "CATALOG_PATH", products)
    monkeypatch.setattr(config, "VENDOR_PRICELISTS_DIR", tmp_path / "vendor")
    monkeypatch.setattr(config, "VENDOR_SUPPLIERS", {"bangkok-christmas": {"label": BC, "book": BC}})
    lookup = tmp_path / "vendor" / "bangkok-christmas" / "cleaned" / "lookup.json"
    lookup.parent.mkdir(parents=True)
    lookup.write_text("{}", encoding="utf-8")
    catalog.refresh()
    vendor_lookup.refresh()
    yield tmp_path
    catalog.refresh()
    vendor_lookup.refresh()


@pytest.fixture
def local(monkeypatch):
    monkeypatch.setattr(settings, "is_local", lambda request: True)


def seed():
    catalog_admin.add_product("BALL-1", "80 mm.", "Glass balls", BC, png_bytes(color=(200, 0, 0, 255)))
    catalog_admin.add_product("TREE-1", "5 Ft.", "Christmas Trees", BC, png_bytes(color=(0, 120, 0, 255)))
    catalog_admin.add_product("MSND-1", "H 215 x D 142 cm", "ต้นคริสต์มาส", MSND,
                              png_bytes(color=(0, 0, 200, 255)), price="7590")


def exported():
    return zipfile.ZipFile(catalog_xlsx.export_zip())


def sheet_rows(zf):
    ws = load_workbook(io.BytesIO(zf.read("products.xlsx")))["สินค้า"]
    header = [c.value for c in ws[1]]
    return header, [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def repack(zf, edit=None, extra_files=None):
    """Rebuild a package from an exported one, optionally editing sheet rows."""
    wb = load_workbook(io.BytesIO(zf.read("products.xlsx")))
    ws = wb["สินค้า"]
    header = [c.value for c in ws[1]]
    if edit:
        edit(ws, header)
    buf = io.BytesIO()
    wb.save(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as new:
        new.writestr("products.xlsx", buf.getvalue())
        for name in zf.namelist():
            if name.startswith("images/"):
                new.writestr(name, zf.read(name))
        for name, data in (extra_files or {}).items():
            new.writestr(name, data)
    out.seek(0)
    return out


def set_cell(ws, header, code, column, value):
    col = header.index(column) + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == code:
            ws.cell(row=row, column=col).value = value
            return
    raise AssertionError(code)


# ---------------------------------------------------------------- export


def test_export_has_every_product_with_its_photo(temp_catalog):
    seed()
    zf = exported()
    header, rows = sheet_rows(zf)

    assert header[:3] == ["รหัสสินค้า", "ชื่อ", "รูป"]
    codes = {r[0]: r for r in rows}
    assert set(codes) == {"BALL-1", "TREE-1", "MSND-1"}
    assert f"images/{codes['BALL-1'][2]}" in zf.namelist()
    assert codes["TREE-1"][header.index("หมวดหมู่หลัก")] == catalog.label_for("tree")
    assert codes["TREE-1"][header.index("ขนาดด้านยาวสุด (มม.)")] == 1524
    assert codes["MSND-1"][header.index("ราคา")] == 7590


# ---------------------------------------------------------------- round trip


def test_an_unedited_export_previews_as_no_change_and_applies_cleanly(temp_catalog):
    seed()
    result = catalog_xlsx.preview(repack(exported()))

    assert result["can_apply"], result["errors"]
    assert result["added"] == [] and result["removed"] == [] and result["changed"] == []
    assert result["unchanged"] == 2
    assert result["skipped_kept_book"] == 1

    applied = catalog_xlsx.apply(result["token"])
    assert applied["imported"] == 2 and applied["kept"] == 1
    assert catalog.find("BALL-1")["size_raw"] == "80 mm."
    assert catalog.find("MSND-1")["price"] == 7590.0  # the kept book is untouched
    assert catalog.is_offered_code("BALL-1")
    assert catalog.image_path("BALL-1").is_file()


def test_edits_show_in_the_preview_and_land_on_apply(temp_catalog):
    seed()

    def edit(ws, header):
        set_cell(ws, header, "BALL-1", "ราคา", 45)
        set_cell(ws, header, "BALL-1", "หมวดหมู่หลัก", "ระฆัง")
        set_cell(ws, header, "BALL-1", "ขนาด", "8 ซม.")
        set_cell(ws, header, "BALL-1", "จำนวน", 6)
        set_cell(ws, header, "BALL-1", "ชื่อ", "ลูกบอลแดง")

    result = catalog_xlsx.preview(repack(exported(), edit))
    assert result["changed"] == [{"code": "BALL-1",
                                  "fields": ["name", "category", "pack_size", "size_raw", "price"]}]

    catalog_xlsx.apply(result["token"])
    row = catalog.find("BALL-1")
    assert catalog.category_of(row) == "bell"  # stated category beats the "ball" section text
    assert catalog.longest_side_mm(row) == 80
    assert row["pack_size"] == 6 and row["name"] == "ลูกบอลแดง"
    assert vendor_lookup.price_for("BALL-1") == 45.0  # prices live in the vendor overlay
    assert "price" not in json.loads(config.CATALOG_PATH.read_text(encoding="utf-8"))[0]


def test_a_row_left_out_is_removed_and_a_new_one_added(temp_catalog):
    seed()
    zf = exported()

    def edit(ws, header):
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == "TREE-1":
                ws.cell(row=row, column=1).value = "TREE-2"

    result = catalog_xlsx.preview(repack(zf, edit))
    assert result["added"] == ["TREE-2"] and result["removed"] == ["TREE-1"]

    catalog_xlsx.apply(result["token"])
    assert catalog.is_offered_code("TREE-2")
    with pytest.raises(ValidationError):
        catalog.find("TREE-1")


def test_colours_come_back_as_colours(temp_catalog):
    seed()
    zf = exported()
    header, rows = sheet_rows(zf)
    ball = next(r for r in rows if r[0] == "BALL-1")[2]

    def edit(ws, header):
        set_cell(ws, header, "BALL-1", "รูป", f"{ball}; green.png")
        set_cell(ws, header, "BALL-1", "ชื่อสี", "แดง; เขียว")

    package = repack(zf, edit, {"images/green.png": png_bytes(color=(0, 200, 0, 255))})
    result = catalog_xlsx.preview(package)
    assert result["can_apply"], result["errors"]
    catalog_xlsx.apply(result["token"])

    variants = catalog.variants_of("BALL-1")
    assert len(variants) == 2
    assert [catalog.colour_name("BALL-1", v) for v in variants] == ["แดง", "เขียว"]
    assert all(catalog.resolve_image_path(v).is_file() for v in variants)


# ---------------------------------------------------------------- refusals


def test_bad_rows_are_errors_and_nothing_is_written(temp_catalog):
    seed()
    before = config.CATALOG_PATH.read_text(encoding="utf-8")

    def edit(ws, header):
        set_cell(ws, header, "BALL-1", "หมวดหมู่หลัก", "ไม่มีหมวดนี้")
        set_cell(ws, header, "BALL-1", "จำนวน", "สองกล่อง")
        set_cell(ws, header, "TREE-1", "รูป", "missing.png")
        ws.append(["BALL-1", "ซ้ำ"])
        ws.append([None, "ไม่มีรหัส"])

    result = catalog_xlsx.preview(repack(exported(), edit))
    assert not result["can_apply"]
    text = " | ".join(result["errors"])
    for expected in ("ไม่รู้จักหมวด", "จำนวน", "missing.png", "รหัสซ้ำ", "ไม่มีรหัสสินค้า"):
        assert expected in text
    assert config.CATALOG_PATH.read_text(encoding="utf-8") == before


def test_a_zip_that_escapes_its_folder_is_refused(temp_catalog):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("../evil.txt", "x")
    out.seek(0)

    with pytest.raises(ValidationError, match="ไม่ปลอดภัย"):
        catalog_xlsx.preview(out)
    assert not (temp_catalog / "data" / "evil.txt").exists()


def test_apply_needs_a_real_preview_token(temp_catalog):
    with pytest.raises(ValidationError):
        catalog_xlsx.apply("../../etc")
    with pytest.raises(ValidationError):
        catalog_xlsx.apply("0" * 32)


def test_apply_backs_up_first(temp_catalog):
    seed()
    shop_overlay.set_fields("BALL-1", {"section": "hand edit"}, speaks_for=("section",))
    catalog.refresh()
    result = catalog_xlsx.preview(repack(exported()))
    applied = catalog_xlsx.apply(result["token"])

    backup = config.DATA_DIR / "backups" / applied["backup"]
    with zipfile.ZipFile(backup) as zf:
        assert {"catalog/products.json", "catalog/product_images.json",
                "data/shop_overlay.json"} <= set(zf.namelist())
        assert "hand edit" in zf.read("data/shop_overlay.json").decode("utf-8")
    # the sheet now owns the section, so the old overlay opinion is gone
    assert "section" not in shop_overlay.fields_for("BALL-1")
    assert catalog.find("BALL-1")["section"] == "hand edit"  # it was exported, so it round-trips


# ---------------------------------------------------------------- endpoints


def test_export_and_import_over_http(client, temp_catalog, local):
    seed()
    response = client.get("/api/catalog/export")
    assert response.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(response.content))

    preview = client.post(
        "/api/catalog/import/preview",
        files={"package": ("catalog.zip", repack(zf).getvalue(), "application/zip")},
    )
    assert preview.status_code == 200, preview.text
    applied = client.post(f"/api/catalog/import/apply/{preview.json()['token']}")
    assert applied.status_code == 200, applied.text
    assert applied.json()["imported"] == 2


def test_a_base_row_missing_blank_fields_still_searches(client, temp_catalog):
    seed()
    path = config.CATALOG_PATH
    rows = json.loads(path.read_text(encoding="utf-8"))
    for key in ("size_raw", "size", "section"):
        rows[0].pop(key, None)
    path.write_text(json.dumps(rows), encoding="utf-8")
    catalog.refresh()

    response = client.get("/api/catalog/search?q=BALL-1")

    assert response.status_code == 200
    assert response.json()["results"][0]["size_raw"] is None


def test_import_is_local_only(client, temp_catalog):
    assert client.get("/api/catalog/export").status_code == 403


# ---------------------------------------------------------------- sizes a person types


@pytest.mark.parametrize("raw, longest", [
    ("30 ซม.", 300), ("8 นิ้ว", 203), ("6 ฟุต", 1829), ("1.5 เมตร", 1500), ("120 มม", 120),
    ("40x60 ซม", 600), ("20 x 30 x 10 cm", 300), ("5 feet", 1524), ("2 m", 2000),
    ("H 215 x D 142 cm", 2150),
])
def test_sizes_in_thai_and_english_units(raw, longest):
    assert catalog.longest_side_mm({"size": catalog.parse_size(raw)}) == longest


def test_a_size_nobody_can_read_stays_unknown():
    assert catalog.parse_size("3 เม็ด") is None

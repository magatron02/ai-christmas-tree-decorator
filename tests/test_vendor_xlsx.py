"""Vendor price list through Excel (backend/services/vendor_xlsx.py): the price import that
works on an installed copy. Same isolated temp lookup as test_vendor_admin.py."""

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from backend import config
from backend.services import settings, vendor_lookup, vendor_overlay, vendor_xlsx
from backend.validation import ValidationError

from test_vendor_admin import set_lookup, temp_catalog, temp_vendor_lookup  # noqa: F401

SUPPLIER = "bangkok-christmas"


@pytest.fixture
def local(monkeypatch):
    monkeypatch.setattr(settings, "is_local", lambda request: True)


@pytest.fixture(autouse=True)
def clean_overlay():
    path = vendor_overlay.overlay_path()
    path.unlink(missing_ok=True)
    vendor_overlay.refresh()
    yield
    path.unlink(missing_ok=True)
    vendor_overlay.refresh()


def edited(rows):
    """The exported sheet with `rows` ({code: {column index: value}}) changed."""
    workbook = load_workbook(io.BytesIO(vendor_xlsx.export_xlsx(SUPPLIER)))
    sheet = workbook[vendor_xlsx.SHEET]
    for row in sheet.iter_rows(min_row=2):
        for index, value in rows.get(row[0].value, {}).items():
            row[index].value = value
    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)
    return out


def test_export_lists_the_suppliers_codes(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0, "name": "บอลแดง", "pack_qty": 12}})

    sheet = load_workbook(io.BytesIO(vendor_xlsx.export_xlsx(SUPPLIER)))[vendor_xlsx.SHEET]

    assert [c.value for c in sheet[2]][:3] == ["071-11", "บอลแดง", 89.0]


def test_an_untouched_export_changes_nothing(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0, "name": "บอลแดง"}})

    result = vendor_xlsx.preview(edited({}), SUPPLIER)

    assert result["changed"] == 0 and result["errors"] == [] and not result["can_apply"]


def test_a_new_price_is_previewed_then_written_to_the_overlay(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}, "071-12": {"price": 45.0}})

    result = vendor_xlsx.preview(edited({"071-11": {2: 95}}), SUPPLIER)
    assert result["changed"] == 1
    assert result["listed"][0]["changes"] == {"price": [89.0, 95.0]}
    assert vendor_lookup.price_for("071-11") == 89.0  # preview writes nothing

    vendor_xlsx.apply(result["token"])

    assert vendor_lookup.price_for("071-11") == 95.0
    assert vendor_lookup.price_for("071-12") == 45.0
    assert json.loads(temp_vendor_lookup.read_text(encoding="utf-8"))["071-11"]["price"] == 89.0


def test_a_price_typed_back_to_the_suppliers_own_leaves_no_opinion(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}})
    vendor_overlay.set_fields("071-11", {"price": 95.0}, speaks_for=("price",))
    vendor_lookup.refresh()

    result = vendor_xlsx.preview(edited({"071-11": {2: 89}}), SUPPLIER)
    vendor_xlsx.apply(result["token"])

    assert "price" not in vendor_overlay.fields_for("071-11")
    assert vendor_lookup.price_for("071-11") == 89.0


def test_other_codes_opinions_survive_an_apply(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}, "071-12": {"price": 45.0, "pack_qty": 6}})
    vendor_overlay.set_fields("071-12", {"pack_qty": 10}, speaks_for=("pack_qty",))
    vendor_lookup.refresh()

    result = vendor_xlsx.preview(edited({"071-11": {2: 95}}), SUPPLIER)
    vendor_xlsx.apply(result["token"])

    assert vendor_overlay.fields_for("071-12") == {"pack_qty": 10}


def test_a_blank_cell_is_no_change(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}})

    result = vendor_xlsx.preview(edited({"071-11": {2: None}}), SUPPLIER)

    assert result["changed"] == 0


def test_a_bad_number_blocks_the_apply(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}})

    result = vendor_xlsx.preview(edited({"071-11": {2: "abc"}}), SUPPLIER)

    assert result["errors"] and not result["can_apply"]


def test_an_unknown_code_is_an_error(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = vendor_xlsx.SHEET
    sheet.append([label for label, _ in vendor_xlsx.COLUMNS])
    sheet.append(["NOPE-1", None, 10, None, None, None])
    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)

    result = vendor_xlsx.preview(out, SUPPLIER)

    assert result["errors"] and not result["can_apply"]


def test_a_pack_quantity_change_is_applied(temp_vendor_lookup):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0, "pack_qty": 12, "pack_unit": "กล่อง"}})

    result = vendor_xlsx.preview(edited({"071-11": {4: 6}}), SUPPLIER)
    vendor_xlsx.apply(result["token"])

    assert vendor_lookup.pack_for("071-11") == {"qty": 6, "unit": "กล่อง"}


def test_apply_rejects_a_made_up_token(temp_vendor_lookup):
    with pytest.raises(ValidationError):
        vendor_xlsx.apply("../../etc/passwd")
    with pytest.raises(ValidationError):
        vendor_xlsx.apply("0" * 32)


def test_http_round_trip_and_local_only(client, temp_vendor_lookup, local, monkeypatch):
    set_lookup(temp_vendor_lookup, {"071-11": {"price": 89.0}})

    exported = client.get(f"/api/vendor/export?supplier={SUPPLIER}")
    assert exported.status_code == 200

    preview = client.post(
        "/api/vendor/import-xlsx/preview", data={"supplier": SUPPLIER},
        files={"package": ("p.xlsx", edited({"071-11": {2: 99}}).read())},
    ).json()
    assert preview["changed"] == 1
    assert client.post(f"/api/vendor/import-xlsx/apply/{preview['token']}").status_code == 200
    assert vendor_lookup.price_for("071-11") == 99.0

    monkeypatch.setattr(settings, "is_local", lambda request: False)
    assert client.get(f"/api/vendor/export?supplier={SUPPLIER}").status_code == 403
